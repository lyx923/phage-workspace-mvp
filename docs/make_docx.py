# make_docx.py
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = (0x1B, 0x6C, 0xA8)

def f(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
    if rFonts.getparent() is None: rPr.append(rFonts)
    for k in ("w:eastAsia", "w:ascii", "w:hAnsi"): rFonts.set(qn(k), name)

def h(doc, text, level=1):
    p = doc.add_heading(level=level); r = p.add_run(text)
    sizes = {0: 22, 1: 16, 2: 13.5, 3: 12}
    f(r, size=sizes.get(level, 12), bold=True,
      color=BLUE if level <= 2 else (0x1E, 0x29, 0x3B))
    return p

def p(doc, text, size=10.5, bold=False, align=None, color=None):
    par = doc.add_paragraph()
    if align == "center": par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f(par.add_run(text), size=size, bold=bold, color=color)
    par.paragraph_format.space_after = Pt(6)
    return par

def b(doc, text, level=0):
    par = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    f(par.add_run(text), size=10.5)
    par.paragraph_format.space_after = Pt(3)
    return par

def code(doc, text):
    par = doc.add_paragraph()
    f(par.add_run(text), name="Consolas", size=9)
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.space_after = Pt(6)
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F4F7FA")
    pPr.append(shd)

def tbl(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        f(c.paragraphs[0].add_run(hd), size=10, bold=True, color=(0xFF, 0xFF, 0xFF))
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "1B6CA8")
        tcPr.append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            f(cells[i].paragraphs[0].add_run(str(v)), size=9.5)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows: r.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

def build(out_path):
    doc = Document()
    doc.styles["Normal"].font.name = "微软雅黑"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.2)

    # ===== 标题 =====
    h(doc, "Phage Intelligence Workspace - MVP v0.7-EXP", 0)
    p(doc, "噬菌体精准治疗 + 竞争情报 + 实验集成\n—— 三域一体的 Ontology 驱动决策智能系统", size=12, bold=True, align="center", color=BLUE)

    # ===== 一、项目简介 =====
    h(doc, "一、项目简介", 1)
    p(doc, "Phage Intelligence Workspace 是一个基于知识图谱和大语言模型的企业级 Ontology 驱动决策智能平台。平台覆盖三个相互连接的领域：")
    b(doc, "Scientific Domain：噬菌体-宿主互作、裂解实验、临床病例、证据包")
    b(doc, "Competitive Intelligence (CI) Domain：竞争对手、研发项目、情报事件、专利、临床试验")
    b(doc, "Engineering Intelligence Domain：工程化噬菌体、工程策略、技术主张、实验结果")
    p(doc, "共享 Foundation 层（SourceArtifact / Pathogen / Organization / Review / AuditEvent / OntologyModule），并通过 Decision & Learning Layer 将外部情报连接成可追溯、可审核、可学习的组织知识系统。")
    p(doc, "当前版本为 v0.7-EXP（实验集成验证版），已在底层数据模型（Graph Schema v2）中全面引入酶基因、实验批次、生物膜实验等实体，实现真实实验数据的结构化接收。")

    h(doc, "1.1 核心验证目标", 2)
    tbl(doc, ["目标", "说明", "状态"], [
        ["V1", "Ontology 能否承载真实临床数据？", "✅ 必填字段填充率 100%"],
        ["V2", "历史病例知识能否被复用？", "✅ 跨病例复用验证通过"],
        ["V3", "实验数据能否结构化录入？", "✅ 酶基因/工程化/生物膜 (v0.7)"],
        ["CI-V1", "竞争对手情报能否统一映射？", "✅ 组织/项目/事件闭环"],
        ["CI-V2", "工程化噬菌体策略能否标准化？", "✅ 策略/构建体关联"],
    ], widths=[2.5, 7.0, 6.5])

    h(doc, "1.2 核心能力", 2)
    for title, items in [
        ("三域知识图谱", [
            "科学域：噬菌体、宿主菌株、裂解实验、临床病例、黄金规则",
            "CI 域：组织、研发项目、情报事件、专利族、临床试验、竞争评估",
            "工程情报域：工程策略、构建体、技术主张、技术结果、技术评估"]),
        ("Ontology 驱动（Palantir-style）", [
            "事实层 / 分析层 / 决策层严格分层",
            "所有事实对象 100% 关联 SourceArtifact",
            "所有状态变更 100% 写入 AuditEvent",
            "所有关键判断须经 Review 审核"]),
        ("双引擎推荐", [
            "规则引擎：确定性逻辑，0 成本，0 幻觉",
            "LLM 引擎：DeepSeek 仅负责文本组织，不产生未检索事实"]),
        ("实验集成闭环 (v0.7 新增)", [
            "实验批次（ExperimentBatch）作为数据血缘锚点",
            "结构化记录酶筛选、工程化改造、生物膜验证数据"]),
    ]:
        p(doc, title, bold=True)
        for it in items: b(doc, it)

    # ===== 二、数据模型 =====
    doc.add_page_break()
    h(doc, "二、数据模型", 1)

    h(doc, "2.1 Foundation 层（跨域共享）", 2)
    tbl(doc, ["实体", "说明", "关键属性"], [
        ["Pathogen", "病原菌主数据", "pathogen_id, species, taxonomy_id, aliases"],
        ["Organization", "组织主数据", "organization_id, canonical_name, organization_type, headquarters_country"],
        ["SourceArtifact", "来源制品", "source_id, source_type, title, url, credibility_tier"],
        ["Review", "统一审核记录", "review_id, review_type, decision, reviewer_id"],
        ["AuditEvent", "不可变审计事件", "audit_id, action_type, object_type, object_id, actor_id, delta, timestamp"],
        ["SchemaVersion", "Schema 版本记录", "version, applied_at, description"],
    ], widths=[4.5, 4.5, 7.0])

    h(doc, "2.2 Scientific Domain (含 v0.7 新增)", 2)
    tbl(doc, ["实体", "说明", "关键属性"], [
        ["Phage", "噬菌体", "phage_id, name, family, is_engineered, engineering_method"],
        ["EnzymeGene", "多糖水解酶基因", "id, name, gene_symbol, specific_activity"],
        ["ExperimentBatch", "实验批次", "id, name, experiment_type, status"],
        ["BiofilmAssay", "生物膜消除实验", "id, control_type, biofilm_reduction_pct"],
        ["LysisAssay", "裂解实验", "assay_id, result, eop_value, od_timeseries"],
        ["ClinicalCase", "临床病例", "case_id, infection_type, phage_treatment, clinical_outcome"],
    ], widths=[4.5, 4.5, 7.0])

    h(doc, "2.3 CI 与决策支持", 2)
    tbl(doc, ["实体", "说明", "关键属性"], [
        ["DevelopmentProgram", "研发项目", "program_id, canonical_name, development_stage"],
        ["IntelligenceEvent", "情报事件", "event_id, event_type, title, event_date"],
        ["IntelligenceProduct", "情报产品", "brief_id, brief_type, review_status"],
        ["DecisionRecord", "决策记录", "decision_id, decision_type, decision_summary"],
        ["EngineeringStrategy", "工程策略", "strategy_id, strategy_type, evidence_maturity"],
    ], widths=[4.5, 4.5, 7.0])

    # ===== 三、快速开始 =====
    doc.add_page_break()
    h(doc, "三、快速开始", 1)

    h(doc, "3.1 环境要求", 2)
    tbl(doc, ["组件", "版本要求"], [
        ["Python", "3.11+"],
        ["Neo4j", "5.x（Community Edition）"],
        ["Redis", "7.x（可选，用于监控缓存）"],
        ["内存", "≥ 8GB"],
    ], widths=[4.0, 8.0])

    h(doc, "3.2 克隆与安装", 2)
    code(doc, "git clone https://github.com/lyx923/phage-workspace-mvp.git\ncd phage-workspace-mvp\npip install -r requirements.txt")

    h(doc, "3.3 启动 Neo4j 与 Redis", 2)
    code(doc,
        "docker run -d --name neo4j -p 7474:7474 -p 7687:7687 \\\n"
        "  -e NEO4J_AUTH=neo4j/your-password neo4j:5.26\n\n"
        "docker run -d --name redis -p 6379:6379 redis:7-alpine")

    h(doc, "3.4 配置 .env", 2)
    code(doc,
        "NEO4J_URI=bolt://localhost:7687\n"
        "NEO4J_USER=neo4j\n"
        "NEO4J_PASSWORD=your-password\n"
        "DS_API_KEY=sk-your-deepseek-api-key\n"
        "DS_MODEL=deepseek-chat")

    h(doc, "3.5 初始化数据", 2)
    p(doc, "启动系统后，进入「系统管理 -> 数据导入」页面，点击「清空并重新导入全部数据」按钮，系统将自动重建约束并注入演示数据。")

    h(doc, "3.6 启动前端", 2)
    code(doc, "python app.py\n# 在浏览器中访问 http://localhost:8501")

    doc.add_page_break()

    # ===== 四、核心设计原则 =====
    h(doc, "四、核心设计原则", 1)
    tbl(doc, ["原则", "落地方式"], [
        ["P1: AI 从真实业务开始", "使用真实历史病例（已去标识化），不使用合成数据"],
        ["P2: AI 学习经过验证的知识", "所有输入数据均经微生物实验室验证"],
        ["P3: AI 放大不替代", "领域专家对输出做最终评审，LLM 仅做文本组织"],
        ["P4: Evidence-driven", "输出标注证据来源和等级（L1-L5），每条推荐可追溯"],
        ["P5: 学习是终点", "知识策展闭环：新病例 → 新知识 → 更好推荐"],
        ["P6: 人工审核不伪造", "系统从不自动批准升级或确认复用"],
    ], widths=[5.0, 10.0])

    # ===== 五、技术栈 =====
    h(doc, "五、技术栈", 1)
    tbl(doc, ["层", "技术"], [
        ["数据存储", "Neo4j 5.x + Redis 7.x"],
        ["LLM", "DeepSeek Chat（OpenAI 兼容 SDK）"],
        ["后端", "Python 3.11+"],
        ["前端", "Dash + dash-bootstrap-components + dash-ag-grid + dash-cytoscape"],
        ["数据处理", "pandas / numpy / scikit-learn"],
    ], widths=[4.0, 11.0])

    # ===== 六、版本历史 =====
    h(doc, "六、版本历史", 1)
    tbl(doc, ["版本", "日期", "主要变更"], [
        ["v0.5", "2026-08-24", "CI Demo 修复；Learning Engine 闭环目标"],
        ["v0.6", "2026-08-31", "Foundation 稳定版；修复归因链语义污染；完成 Dash 重构"],
        ["v0.7-EXP", "2026-10-29", "实验集成验证；新增 EnzymeGene / ExperimentBatch / BiofilmAssay；引入 SchemaVersion 迁移机制"],
    ], widths=[2.5, 3.0, 9.5])

    # ===== 尾注 =====
    p(doc, "")
    p(doc, "Phage Intelligence Workspace — 让上一个病例的经验，成为下一个病例的起点；让外部情报，成为内部决策的依据。",
      size=11, bold=True, align="center", color=BLUE)

    doc.save(out_path)
    print(f"✅ 已生成：{out_path}")


if __name__ == "__main__":
    root = os.path.dirname(os.path.abspath(__file__))
    out_dir = os.path.join(root, "")
    os.makedirs(out_dir, exist_ok=True)
    build(os.path.join(out_dir, "Phage_Intelligence_Workspace_v0.7.docx"))